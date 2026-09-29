import os
import pymupdf
from pathlib import Path
from docx import Document
from docx.shared import Inches

ruta_descargas = Path.home()


def encontrar_zona_por_lineas(pagina, ancho_minimo, tolerancia=2):
    lineas_horizontales = []
    lineas_verticales = []

    for dibujo in pagina.get_drawings():
        for item in dibujo["items"]:
            tipo = item[0]

            if tipo == "l":
                p1, p2 = item[1], item[2]
                x0, x1 = sorted([p1.x, p2.x])
                y0, y1 = sorted([p1.y, p2.y])

            elif tipo == "re":
                r = item[1]
                x0, x1 = r.x0, r.x1
                y0, y1 = r.y0, r.y1

            else:
                continue

            ancho = x1 - x0
            alto = y1 - y0

            if ancho > ancho_minimo and alto <= tolerancia:
                lineas_horizontales.append((x0, x1, y0, ancho))

            elif alto > tolerancia and ancho <= tolerancia:
                lineas_verticales.append((x0, y0, y1, alto))

    if not lineas_horizontales:
        return None

    x0, x1, y_inicio, _ = max(
        lineas_horizontales,
        key=lambda l: l[3]
    )

    if lineas_verticales:
        _, _, _, altura = max(
            lineas_verticales,
            key=lambda l: l[3]
        )
        y_fin = y_inicio + altura
    else:
        y_fin = pagina.rect.y1

    return pymupdf.Rect(
        x0,
        y_inicio,
        x1,
        y_fin
    )


def detectar_zona(pdf_origen):
    doc = pymupdf.open(pdf_origen)

    try:
        pagina = doc[0]

        UMBRAL_RELATIVO = 0.89
        MIN_ALTURA_RECTANGULO = 5
        MARGEN_INFERIOR = 70

        ancho_minimo = pagina.rect.width * UMBRAL_RELATIVO
        rectangulos = []

        for dibujo in pagina.get_drawings():
            rect = dibujo["rect"]

            if (
                rect.width > ancho_minimo
                and rect.height > MIN_ALTURA_RECTANGULO
            ):
                rectangulos.append(rect)

        zona = None

        if rectangulos:
            rectangulo_mayor = max(
                rectangulos,
                key=lambda r: r.width * r.height
            )

            zona = pymupdf.Rect(rectangulo_mayor)
            zona.y1 -= MARGEN_INFERIOR

        else:
            zona = encontrar_zona_por_lineas(
                pagina,
                ancho_minimo
            )

            if zona is not None:
                zona.y1 -= MARGEN_INFERIOR

        if zona is None:
            return None

        zona &= pagina.rect

        if zona.is_empty:
            return None

        if zona.width <= 0 or zona.height <= 0:
            return None

        return zona

    finally:
        doc.close()


def crear_vista_previa(pdf_origen, salida):
    doc = pymupdf.open(pdf_origen)

    try:
        pagina = doc[0]
        zona = detectar_zona(pdf_origen)

        if zona is None:
            return False

        pix = pagina.get_pixmap(
            clip=zona,
            matrix=pymupdf.Matrix(2, 2),
            alpha=False
        )

        pix.save(salida)

        return True

    finally:
        doc.close()


def obtener_info_pdf(pdf_origen):
    doc = pymupdf.open(pdf_origen)

    try:
        pagina = doc[0]

        tamaño = os.path.getsize(pdf_origen)
        tamaño_mb = tamaño / (1024 * 1024)

        texto = pagina.get_text()

        return {
            "nombre": os.path.basename(pdf_origen),
            "ruta": pdf_origen,
            "paginas": len(doc),
            "tamaño": f"{tamaño_mb:.2f} MB",
            "ancho": round(pagina.rect.width, 2),
            "alto": round(pagina.rect.height, 2),
            "texto": texto
        }

    finally:
        doc.close()

def table_order_list_word(lista, docx_path):
    doc_word = Document(docx_path) if os.path.exists(docx_path) else Document()

    filas = (len(lista) + 1) // 2
    table = doc_word.add_table(rows=filas, cols=2)

    for i, archivo in enumerate(lista):
        doc_pdf = pymupdf.open(archivo)

        try:
            pagina = doc_pdf[0]
            zona = detectar_zona(archivo)

            if zona is None:
                raise ValueError(
                    f"No se encontró una zona válida en '{archivo}'."
                )

            pix = pagina.get_pixmap(
                clip=zona,
                matrix=pymupdf.Matrix(2, 2),
                alpha=False
            )

            imagen_temporal = os.path.join(
                str(ruta_descargas),
                f"recorte_temp_{i}.png"
            )

            pix.save(imagen_temporal)

        finally:
            doc_pdf.close()

        fila = i // 2
        columna = i % 2

        cell = table.cell(fila, columna)
        paragraph = cell.paragraphs[0]
        run = paragraph.add_run()
        run.add_picture(imagen_temporal, width=Inches(3))

        os.remove(imagen_temporal)

    doc_word.save(docx_path)

    return f"Proceso completado: {docx_path}"
